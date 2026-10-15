"""Checks the explorer data against the published source and flow identities."""
import hashlib
import json
import math
import re
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = ROOT.parent / "MFA_V1.4.1-G-0426" / "04-MFA" / "outputs"
DATA = json.loads((ROOT / "data" / "results.json").read_text(encoding="utf-8"))
MANIFEST = json.loads((ROOT / "data" / "manifest.json").read_text(encoding="utf-8"))


class ExplorerDataChecks(unittest.TestCase):
    def test_release_coverage_and_hash(self):
        self.assertEqual(DATA["release"]["status"], "published")
        self.assertEqual(DATA["release"]["model_year"], 2020)
        self.assertEqual(len(DATA["entities"]), 246)
        self.assertEqual(len(DATA["groups"]), 54)
        self.assertEqual(len(DATA["metrics"]), 203)
        self.assertEqual(len(DATA["statistics"]), 6)
        self.assertEqual(hashlib.sha256((ROOT / "data" / "results.json").read_bytes()).hexdigest(), MANIFEST["output"]["sha256"])
        self.assertEqual({e["id"] for e in DATA["entities"]}, set(DATA["national"]))
        self.assertEqual({g["id"] for g in DATA["groups"]}, set(DATA["aggregated"]))

    def test_absent_coefficient_differs_from_zero(self):
        aruba = DATA["national"]["ABW"]
        self.assertIsNone(aruba["coefficients"]["C8"]["Mean"])
        self.assertEqual(aruba["processes"]["P10"]["Mean"], 0)
        self.assertEqual(aruba["processes"]["P17"]["Mean"], 0)
        coefficients = [stats for r in DATA["national"].values() for stats in r["coefficients"].values()]
        self.assertEqual(sum(v is None for stats in coefficients for v in stats.values()), 5676)

    def test_source_labels_and_display_units_are_distinct(self):
        p1aa = next(m for m in DATA["metrics"] if m["family"] == "processes" and m["id"] == "P1aa")
        self.assertEqual(p1aa["unit"], "t·y-1")
        self.assertEqual(p1aa["display_unit"], "tonnes/year")
        self.assertIsNone(p1aa["unit_note"])
        self.assertEqual(next(e for e in DATA["entities"] if e["id"] == "XKO")["name"], "Kosovo")
        self.assertTrue(any(g["scheme"] == "UN Region" and g["name"] == "NA" for g in DATA["groups"]))

    def test_mean_flows_and_plastic_decomposition(self):
        branches = {}
        for link in DATA["flow_links"]:
            branches.setdefault(link["source"], []).append(link["metric_id"])
        identities = {
            "plas_litter_em": ["P2a"],
            "plas_uncol_em": ["P20a", "P21a"],
            "plas_collection_em": ["P6a"],
            "plas_disp_em": ["P19a", "P17a"],
            "plas_recy_em": ["P31aa", "P31ab", "P33aa", "P33ab", "P30aa", "P30ab", "P32aa", "P32ab"],
            "plas_debris_em": ["P2a", "P20a", "P6a", "P19a", "P31aa", "P31ab", "P33aa", "P33ab"],
            "plas_burn_em": ["P21a", "P17a", "P30aa", "P30ab", "P32aa", "P32ab"],
        }
        for scope in ["national", "aggregated"]:
            for record_id, record in DATA[scope].items():
                process, output = record["processes"], record["outputs"]
                for parent, children in branches.items():
                    self.assertTrue(math.isclose(process[parent]["Mean"], sum(process[c]["Mean"] for c in children), rel_tol=1e-11, abs_tol=1e-7), (record_id, parent))
                for metric, children in identities.items():
                    self.assertTrue(math.isclose(output[metric]["Mean"], sum(process[c]["Mean"] for c in children), rel_tol=1e-11, abs_tol=1e-7), (record_id, metric))


@unittest.skipUnless(openpyxl and SOURCE_ROOT.exists(), "Optional source workbooks/openpyxl are not available")
class SourceValueChecks(unittest.TestCase):
    def test_every_exported_source_value_and_metadata_field(self):
        for kind, scope in [("national", "national"), ("aggregated", "aggregated")]:
            source = MANIFEST["sources"][kind]
            path = SOURCE_ROOT / source["filename"]
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), source["sha256"])
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                native_info = [{"row": index, "label": row[1], "value": row[2]}
                               for index, row in enumerate(book["00_Info"].iter_rows(values_only=True), start=1)
                               if len(row) > 2 and (row[1] is not None or row[2] is not None)]
                self.assertEqual(source["native_info_sheet"], native_info)
                descriptions = {r["id"]: r for r in DATA["entities" if kind == "national" else "groups"]}
                for family, sheet in [("processes", "Processes"), ("coefficients", "Coefficients"), ("outputs", "Outputs")]:
                    rows = book[sheet].iter_rows(values_only=True)
                    headers = next(rows)
                    offset = 10 if kind == "national" else 5
                    for row in rows:
                        if row[0] is None:
                            continue
                        native = dict(zip(headers, row))
                        if kind == "national":
                            key = native["ISO3"]
                        else:
                            slug = lambda value: re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
                            key = f"{slug(native['Aggregation_lv1'])}:{slug(native['Aggregation_lv2'])}"
                        metadata_fields = {"population": "Population_2020", "rural_share": "Rural_share"}
                        if kind == "national":
                            metadata_fields.update({"name": "Country", "income": "Income_cat", "region": "Region", "subregion": "Subregion", "intermediate_region": "Intermediate_region", "oecd": "OECD_region"})
                        else:
                            metadata_fields.update({"name": "Aggregation_lv2", "scheme": "Aggregation_lv1"})
                        for exported_field, native_field in metadata_fields.items():
                            self.assertEqual(descriptions[key][exported_field], native[native_field], (key, exported_field))
                        exported = DATA[scope][key][family]
                        for field in headers[offset:]:
                            self.assertEqual(exported[field][native["Statistic"]], native[field], (key, family, field, native["Statistic"]))
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
