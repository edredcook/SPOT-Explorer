"""Extract the published SPOT national and grouped results for the explorer.

Run with the bundled Python or another Python with openpyxl installed:
    python tools/prepare_data.py

The original workbooks remain outside this application. Source values are not
rounded, rescaled, filled or recalculated. Workbook provenance and extraction
checks are written to data/manifest.json alongside the compact results.
"""
from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import math
import re
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


APP_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = APP_ROOT.parent / "MFA_V1.4.1-G-0426"
VERSION = "V1.4.1-G-0426"
STATISTICS = ["Mean", "5th percentile", "Lower quartile", "Median", "Upper quartile", "95th percentile"]
FAMILIES = {
    "processes": ("Processes", "01a_Processes_legend"),
    "coefficients": ("Coefficients", "02a_Coefficients_legend"),
    "outputs": ("Outputs", "03a_Outputs_legend"),
}
LINKS = {
    "dataset": "https://doi.org/10.5281/zenodo.22231994",
    "methods": "https://doi.org/10.1038/s41586-024-07758-6",
    "code": "https://github.com/SPOT-Model/SPOT-MFA",
}

# The full MSW system, using the same process equations and annual masses as the
# source. A link's value is the target process, not a second derived estimate.
BRANCHES = {
    "P1": ["P2", "P3"],
    "P3": ["P4", "P5"],
    "P4": ["P20", "P21"],
    "P5": ["P6", "P7"],
    "P7": ["P8", "P9"],
    "P8": ["P10", "P11"],
    "P9": ["P12i", "P12ii", "P13"],
    "P10": ["P16", "P17"],
    "P16": ["P18", "P19"],
    "P13": ["P14", "P15"],
}
NODE_LABELS = {
    "P1": "Waste generated", "P2": "Uncollected litter",
    "P3": "Waste excluding litter", "P4": "Uncollected waste",
    "P5": "Collected waste", "P6": "Collection losses",
    "P7": "Collected and delivered", "P8": "Sent for disposal",
    "P9": "Treatment and recovery", "P10": "Uncontrolled disposal",
    "P11": "Controlled disposal", "P12i": "Composting and anaerobic digestion",
    "P12ii": "Incineration", "P13": "Collected for dry recycling",
    "P14": "Informal recycling collection", "P15": "Formal recycling collection",
    "P16": "Uncontrolled disposal, unburned", "P17": "Disposal openly burned",
    "P18": "Unburned and retained", "P19": "Disposal debris losses",
    "P20": "Uncollected and unburned", "P21": "Uncollected and openly burned",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def numeric(value, context: str):
    """Preserve a source number or missing cell; reject unexpected text/errors."""
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Expected a numeric source cell at {context}: {value!r}")
    if not math.isfinite(value):
        raise ValueError(f"Nonfinite source value at {context}")
    return value


def source_info(workbook):
    """Retain the native Info sheet, including its historical date labels."""
    sheet = workbook["00_Info"]
    return [
        {"row": index, "label": row[1], "value": row[2]}
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
        if len(row) > 2 and (row[1] is not None or row[2] is not None)
    ]


def display_unit(unit, family, metric_id):
    if family == "coefficients":
        return "%", "The source coefficient definitions describe percentages."
    if family == "processes" and metric_id in {"P1a", "P1aa", "P1ab"} and unit in {"t·y-2", "t·y-3", "t·y-4"}:
        return "tonnes/year", f"Annual mass, consistent with the corresponding generation output. The original process legend reads {unit}. Values are unchanged."
    units = {
        "t·y-1": "tonnes/year", "g·y-1": "grams/year",
        "kg·cap-1·d-1": "kg/person/day", "kg·cap-1·y-1": "kg/person/year",
        "People": "people",
    }
    if unit and unit.startswith("%"):
        return "%", None
    return units.get(unit, unit), None


def read_metrics(workbook):
    metrics = []
    for family, (sheet_name, legend_name) in FAMILIES.items():
        sheet = workbook[sheet_name]
        result_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        result_columns = {name: i + 1 for i, name in enumerate(result_headers)}
        legend = workbook[legend_name]
        headers = next(legend.iter_rows(min_row=1, max_row=1, values_only=True))
        for row_number, row in enumerate(legend.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue
            original = dict(zip(headers, row))
            metric_id, short_name, description = row[:3]
            unit = row[3] if len(row) > 3 else None
            shown_unit, note = display_unit(unit, family, metric_id)
            column = get_column_letter(result_columns[metric_id])
            metric = {
                "id": metric_id, "family": family,
                "name": description, "description": description,
                "unit": unit, "display_unit": shown_unit,
                "unit_note": note,
                "source_sheet": sheet_name, "source_field": metric_id,
                "source_range": f"{column}2:{column}{sheet.max_row}",
                "source_legend_sheet": legend_name,
                "source_legend_range": f"A{row_number}:{get_column_letter(len(headers))}{row_number}",
                "source_short_name": short_name,
                "source_metadata": original,
            }
            if family == "outputs" and metric_id == "WP_no":
                metric["interpretation_note"] = "Modelled informal collector workforce; the release README describes full-time-equivalent collectors."
            if family == "outputs" and metric_id.startswith("exp_"):
                metric["interpretation_note"] = "Modelled exposure count. This is not a measured pollutant concentration or health outcome."
            metrics.append(metric)
    if len({(m["family"], m["id"]) for m in metrics}) != len(metrics):
        raise ValueError("Duplicate metric IDs within a family")
    return metrics


def read_results(workbook, grouped=False):
    records, descriptions, sheet_counts = {}, {}, {}
    group_keys = {}
    for family, (sheet_name, _) in FAMILIES.items():
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        offset = 5 if grouped else 10
        fields = headers[offset:]
        seen = set()
        value_count = null_count = zero_count = 0
        for row_number, values in enumerate(rows, start=2):
            if values[0] is None:
                continue
            raw = dict(zip(headers, values))
            statistic = raw["Statistic"]
            if statistic not in STATISTICS:
                raise ValueError(f"Unexpected statistic {statistic!r}")
            if grouped:
                scheme, name = raw["Aggregation_lv1"], raw["Aggregation_lv2"]
                record_id = f"{slug(scheme)}:{slug(name)}"
                pair = (scheme, name)
                if record_id in group_keys and group_keys[record_id] != pair:
                    raise ValueError(f"Group slug collision: {pair!r}")
                group_keys[record_id] = pair
                metadata = {"id": record_id, "name": name, "scheme": scheme,
                            "population": raw["Population_2020"], "rural_share": raw["Rural_share"]}
            else:
                record_id = raw["ISO3"]
                metadata = {
                    "id": record_id, "name": raw["Country"],
                    "population": raw["Population_2020"], "rural_share": raw["Rural_share"],
                    "income": raw["Income_cat"], "region": raw["Region"],
                    "subregion": raw["Subregion"], "intermediate_region": raw["Intermediate_region"],
                    "oecd": raw["OECD_region"],
                }
            key = (record_id, statistic)
            if key in seen:
                raise ValueError(f"Duplicate source row: {sheet_name} {key!r}")
            seen.add(key)
            if record_id in descriptions and descriptions[record_id] != metadata:
                raise ValueError(f"Metadata differs between source rows: {record_id}")
            descriptions[record_id] = metadata
            target = records.setdefault(record_id, {}).setdefault(family, {})
            for column_number, field in enumerate(fields, start=offset + 1):
                value = numeric(raw[field], f"{sheet_name}!{get_column_letter(column_number)}{row_number}")
                target.setdefault(field, {})[statistic] = value
                value_count += 1
                null_count += value is None
                zero_count += value is not None and value == 0
        sheet_counts[family] = {"sheet": sheet_name, "rows": len(seen), "metrics": len(fields),
                                "values_including_null": value_count, "nulls": null_count, "zeros": zero_count}
    for record_id, families in records.items():
        for family, values in families.items():
            for field, stats in values.items():
                if set(stats) != set(STATISTICS):
                    raise ValueError(f"Missing statistics for {record_id}/{family}/{field}")
                ordered = [stats[q] for q in STATISTICS[1:] if stats[q] is not None]
                if ordered != sorted(ordered):
                    raise ValueError(f"Quantile order error for {record_id}/{family}/{field}")
    return records, list(descriptions.values()), sheet_counts


def flow_metadata(national, aggregated, metrics):
    process_metrics = {m["id"]: m for m in metrics if m["family"] == "processes"}
    node_ids = list(dict.fromkeys(["P1"] + [child for children in BRANCHES.values() for child in children]))
    nodes = [{"id": node_id, "name": NODE_LABELS[node_id], "metric_id": node_id,
              "family": "processes", "description": process_metrics[node_id]["description"]}
             for node_id in node_ids]
    links = [{"source": parent, "target": child, "metric_id": child, "family": "processes"}
             for parent, children in BRANCHES.items() for child in children]
    max_relative_error = 0.0
    for records in (national, aggregated):
        for record_id, record in records.items():
            process = record["processes"]
            for parent, children in BRANCHES.items():
                parent_mass = process[parent]["Mean"]
                child_mass = sum(process[child]["Mean"] for child in children)
                relative_error = abs(parent_mass - child_mass) / max(abs(parent_mass), 1)
                max_relative_error = max(max_relative_error, relative_error)
                if not math.isclose(parent_mass, child_mass, rel_tol=1e-11, abs_tol=1e-7):
                    raise ValueError(f"Flow does not balance: {record_id}/{parent}")
    return nodes, links, max_relative_error


def world_checks(national, aggregated, metrics):
    world = aggregated["world:world"]
    additive = [m for m in metrics if m["family"] == "processes" or
                (m["family"] == "outputs" and m["unit"] in {"t·y-1", "People"})]
    max_relative_error = 0.0
    for metric in additive:
        family, field = metric["family"], metric["id"]
        national_sum = sum(record[family][field]["Mean"] for record in national.values())
        world_mean = world[family][field]["Mean"]
        max_relative_error = max(max_relative_error, abs(world_mean - national_sum) / max(abs(world_mean), 1))
        if not math.isclose(national_sum, world_mean, rel_tol=1e-11, abs_tol=1e-7):
            raise ValueError(f"World mean differs from national sum: {family}/{field}")
    return {"additive_metrics_checked": len(additive), "maximum_relative_residual": max_relative_error}


def write_json(path, payload, compact=False):
    path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False,
                               separators=(",", ":") if compact else None,
                               indent=None if compact else 2) + "\n", encoding="utf-8")


def prepare(source_root: Path, output: Path):
    source_paths = {
        "national": source_root / "04-MFA" / "outputs" / f"SPOT_MFA_National_results_{VERSION}_2000i.xlsx",
        "aggregated": source_root / "04-MFA" / "outputs" / f"SPOT_MFA_Other_aggregation_results_{VERSION}_2000i.xlsx",
    }
    hashes_before = {kind: sha256(path) for kind, path in source_paths.items()}
    books = {kind: openpyxl.load_workbook(path, read_only=True, data_only=True) for kind, path in source_paths.items()}
    try:
        metrics = read_metrics(books["national"])
        national, entities, national_counts = read_results(books["national"])
        aggregated, groups, aggregate_counts = read_results(books["aggregated"], grouped=True)
        nodes, links, flow_error = flow_metadata(national, aggregated, metrics)
        checks = world_checks(national, aggregated, metrics)
        native_metadata = {kind: source_info(book) for kind, book in books.items()}
    finally:
        for book in books.values():
            book.close()
    if hashes_before != {kind: sha256(path) for kind, path in source_paths.items()}:
        raise RuntimeError("A source workbook changed during extraction")
    release = {
        "id": VERSION, "version": VERSION, "title": "SPOT global results",
        "status": "published", "model_year": 2020, "iterations": 2000,
        "name": "Spatio-temporal quantification of plastic pollution origins and transport (SPOT) model",
        "component": "Material flow analysis and emissions inventory",
        "citation_year": 2026,
        "citation": "Cottom, J. W., Cook, E. & Velis, C. A. (2026). SPOT model - MFA Version 1.4.1-G-0426. Zenodo.",
        "source_links": LINKS, "links": LINKS,
        "methodology": "Cottom, Cook & Velis (2024), with the published release's Lancet Countdown methodological updates.",
        "scope": "Municipal solid waste and macroplastic emissions. These results describe 2020 and include modelled open-burning atmospheric emissions and exposure estimates.",
        "uncertainty": "The 5th and 95th percentiles span a central 90% modelled interval. Mean and median are available separately.",
        "license": "CC BY 4.0, subject to source-specific notices for reused third-party material.",
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
        "entity_label": "countries and territories",
        "null_note": "A missing coefficient is unavailable or undefined; it is not zero.",
    }
    payload = {"release": release, "statistics": STATISTICS, "entities": entities,
               "groups": groups, "metrics": metrics, "national": national,
               "aggregated": aggregated, "flow_nodes": nodes, "flow_links": links,
               "flow": {"statistic": "Mean", "unit": "tonnes/year", "family": "processes",
                        "title": "Municipal solid waste flows",
                        "note": "Mean annual MSW mass. Recycling branches show collection for recycling, not final secondary material production. Quantile flows should not be added into a balanced network.",
                        "source": "4b-SystemofEquations.R, full MFA process equations P1-P21."}}
    output.mkdir(parents=True, exist_ok=True)
    write_json(output / "results.json", payload, compact=True)
    (output / "results.json.gz").write_bytes(gzip.compress((output / "results.json").read_bytes(), mtime=0))
    manifest = {
        "schema_version": 1, "release": release,
        "extractor": "tools/prepare_data.py", "sources": {
            kind: {"filename": path.name, "relative_source_path": str(path.relative_to(source_root)).replace('\\', '/'),
                   "bytes": path.stat().st_size, "sha256": hashes_before[kind],
                   "native_info_sheet": native_metadata[kind]}
            for kind, path in source_paths.items()
        },
        "counts": {"entities": len(entities), "groups": len(groups), "statistics": len(STATISTICS),
                   "metrics": len(metrics), "metrics_by_family": dict(Counter(m["family"] for m in metrics)),
                   "national_sheets": national_counts, "aggregate_sheets": aggregate_counts},
        "checks": {"source_hashes_unchanged": True, "duplicate_keys": 0,
                   "all_source_numbers_preserved": True, "quantiles_monotonic": True,
                   "world_mean_aggregation": checks,
                   "flow_branches_checked": len(BRANCHES) * (len(entities) + len(groups)),
                   "flow_maximum_relative_residual": flow_error},
        "transformations": ["Source cell numbers and nulls copied without rounding, scaling or recalculation.",
                            "Source field identifiers and country/territory keys retained.",
                            "Group identifiers formed from scheme and name slugs; original labels retained.",
                            "Readable display-unit labels; original legend units retained without numeric conversion.",
                            "Literal NA grouping labels retained as text."],
        "output": {"filename": "results.json", "bytes": (output / "results.json").stat().st_size,
                   "gzip_filename": "results.json.gz", "gzip_bytes": (output / "results.json.gz").stat().st_size,
                   "sha256": sha256(output / "results.json")},
    }
    write_json(output / "manifest.json", manifest)
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=APP_ROOT / "data")
    arguments = parser.parse_args()
    result = prepare(arguments.source_root.resolve(), arguments.output.resolve())
    print(json.dumps({"counts": result["counts"], "output": result["output"], "checks": result["checks"]}, indent=2))
